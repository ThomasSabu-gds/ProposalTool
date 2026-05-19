"""
Knowledge Base — persistent ChromaDB + per-template TOC sidecar JSON.

Layout on disk:
  datalake/.kb_index/
    chroma/            <- persistent ChromaDB collection
    manifest.json      <- { "<filename>": {"mtime": ..., "size": ...}, ... }
    toc/<slug>.json    <- per-template structured outline (sections, slides, summaries)

Indexing is INCREMENTAL — only re-processes files whose (mtime, size) changed
since the last sync. New / updated / removed templates are handled automatically.

Embedding strategy:
  1. Azure OpenAI text-embedding-3-small  (if OPEN_AI_EMBEDDING_ENDPOINT is set)
  2. Local sentence-transformers          (if package + model available)
  3. SimpleHashEmbedding                  (last-resort word-bag, no semantics)
"""

import os
import re
import json
import time
import hashlib
import threading
from pathlib import Path
from typing import Optional

import numpy as np
import chromadb
from chromadb.api.types import EmbeddingFunction, Documents, Embeddings
from pptx import Presentation


# ─── Embedding backends ──────────────────────────────────────────────────────

class SimpleHashEmbedding(EmbeddingFunction):
    """Word-bag hash. Last-resort, no semantics — kept only as a backup."""

    def __call__(self, input: Documents) -> Embeddings:
        embeddings = []
        for doc in input:
            words = set(re.findall(r"\b[a-z]{3,}\b", doc.lower()))
            vec = np.zeros(384, dtype=np.float32)
            for w in words:
                h = int(hashlib.md5(w.encode()).hexdigest(), 16)
                vec[h % 384] += 1.0
            n = np.linalg.norm(vec)
            if n > 0:
                vec = vec / n
            embeddings.append(vec.tolist())
        return embeddings


class AzureOpenAIEmbedding(EmbeddingFunction):
    """Azure OpenAI text-embedding-3-small. Persistent ChromaDB caches vectors,
    so embedding cost is incurred once per slide."""

    def __init__(self, client, deployment: str, batch_size: int = 16):
        self.client = client
        self.deployment = deployment
        self.batch_size = batch_size

    def __call__(self, input: Documents) -> Embeddings:
        out = []
        for start in range(0, len(input), self.batch_size):
            batch = list(input[start : start + self.batch_size])
            # Azure rejects empty strings — substitute a single space
            cleaned = [(t if t and t.strip() else " ") for t in batch]
            for attempt in range(4):
                try:
                    resp = self.client.embeddings.create(
                        model=self.deployment, input=cleaned
                    )
                    out.extend([d.embedding for d in resp.data])
                    break
                except Exception as exc:
                    if attempt == 3:
                        # Pad with zero vectors so caller doesn't crash; logs the issue
                        print(f"  [KB] embedding batch failed after retries: {exc}")
                        out.extend([[0.0] * 1536 for _ in cleaned])
                        break
                    time.sleep(1.5 * (attempt + 1))
        return out


class LocalSentenceTransformerEmbedding(EmbeddingFunction):
    """sentence-transformers/all-MiniLM-L6-v2 — 384-d, runs on CPU, ~80MB download."""

    def __init__(self):
        from sentence_transformers import SentenceTransformer  # type: ignore
        self.model = SentenceTransformer("all-MiniLM-L6-v2")

    def __call__(self, input: Documents) -> Embeddings:
        vecs = self.model.encode(list(input), normalize_embeddings=True)
        return vecs.tolist()


def select_embedding_fn(embedding_client=None, embedding_deployment=None):
    """Pick the best available embedding backend. Logs the choice once."""
    if embedding_client and embedding_deployment:
        print(f"  [KB] embedding backend: Azure ({embedding_deployment})")
        return AzureOpenAIEmbedding(embedding_client, embedding_deployment)
    try:
        emb = LocalSentenceTransformerEmbedding()
        print("  [KB] embedding backend: local sentence-transformers (MiniLM-L6-v2)")
        return emb
    except Exception:
        pass
    print("  [KB] embedding backend: SimpleHashEmbedding (degraded — no semantic search)")
    return SimpleHashEmbedding()


# ─── Section / sector classification (rule-based shortcuts) ──────────────────

SECTION_TYPES = {
    "cover": ["cover", "title slide"],
    "cover_letter": ["dear", "thank you for inviting", "cover letter"],
    "toc": ["table of content", "agenda", "contents"],
    "firm_profile": ["why ey", "why us", "about ey", "firm profile", "about us"],
    "experience": [
        "experience", "credential", "case stud", "project deliver",
        "we worked", "we have deliver", "track record", "past project",
    ],
    "understanding": [
        "our understanding", "point of view", "market insight",
        "key consideration", "privatization", "sector overview",
    ],
    "methodology": [
        "methodology", "approach", "our approach", "framework",
        "workstream", "phase", "how we will",
    ],
    "team": [
        "our team", "team structure", "key personnel", "organis",
        "resource", "staffing", "cv", "profile",
    ],
    "timeline": ["timeline", "gantt", "work plan", "schedule", "milestone"],
    "pricing": ["pricing", "fee", "financial proposal", "cost", "budget"],
    "appendix": ["appendix", "annex", "supporting", "certificate"],
}


def classify_section(title: str, content: str) -> str:
    combined = (title + " " + content).lower()
    for section_type, keywords in SECTION_TYPES.items():
        for kw in keywords:
            if kw in combined:
                return section_type
    return "other"


def detect_sector(prs_text: str, filename: str) -> str:
    combined = (filename + " " + prs_text[:3000]).lower()
    sectors = {
        "Healthcare": ["healthcare", "hospital", "clinical", "patient", "medical", "health"],
        "Oil & Gas": ["oil", "gas", "fuel", "petroleum", "energy", "refinery", "upstream"],
        "Technology": ["technology", "it ", "digital", "software", "cloud", "cyber", "data"],
        "Infrastructure": ["infrastructure", "highway", "road", "bridge", "construction", "ppp"],
        "Financial Services": ["banking", "financial", "insurance", "fintech", "capital market"],
    }
    for sector, keywords in sectors.items():
        matches = sum(1 for kw in keywords if kw in combined)
        if matches >= 2:
            return sector
    return "General"


# ─── TOC sidecar JSON ────────────────────────────────────────────────────────

TOC_SYSTEM = """You are building a structured outline of a proposal-template deck.

For the slides given below, group consecutive slides into SECTIONS and label
each slide with its purpose. Output strictly valid JSON:

{
  "sections": [
    { "name": "<short section title>",
      "purpose": "cover|toc|exec_summary|understanding|methodology|experience|team|timeline|pricing|firm_profile|appendix|other",
      "slide_range": [<first_idx>, <last_idx>] }
  ],
  "slides": [
    { "index": <0-based>,
      "title": "<short title — empty string ok>",
      "purpose": "title_slide|section_header|content|diagram|table|cv|cover_letter|toc|closing|divider",
      "summary": "<one sentence (<=120 chars) describing what this slide contains in the template>" }
  ]
}

Be terse. Do not invent content. If a slide is purely visual, summary can be
"Visual divider with branded imagery" etc."""


def build_toc_with_llm(slides_meta: list, client, deployment: str) -> dict:
    """Single LLM call producing the TOC + per-slide outline."""
    payload = []
    for s in slides_meta:
        payload.append({
            "index": s["index"],
            "layout": s.get("layout", ""),
            "title": s.get("title", "")[:120],
            "text_preview": (s.get("content", "") or "")[:280],
            "word_count": s.get("word_count", 0),
        })
    try:
        resp = client.chat.completions.create(
            model=deployment,
            messages=[
                {"role": "system", "content": TOC_SYSTEM},
                {"role": "user", "content": json.dumps(payload, default=str)[:14000]},
            ],
            response_format={"type": "json_object"},
            temperature=0.05,
            max_tokens=4096,
        )
        return json.loads(resp.choices[0].message.content)
    except Exception as exc:
        print(f"  [KB] TOC LLM failed, using fallback: {exc}")
        # Heuristic fallback — one section, slides labelled by rule classifier
        return {
            "sections": [
                {"name": "All slides", "purpose": "other",
                 "slide_range": [0, max(0, len(slides_meta) - 1)]}
            ],
            "slides": [
                {
                    "index": s["index"],
                    "title": s.get("title", ""),
                    "purpose": classify_section(s.get("title", ""), s.get("content", "")),
                    "summary": (s.get("content", "") or "")[:120],
                }
                for s in slides_meta
            ],
        }


# ─── Knowledge Base ──────────────────────────────────────────────────────────

KB_DIR_NAME = ".kb_index"
MANIFEST_NAME = "manifest.json"


class KnowledgeBase:
    """Persistent ChromaDB-backed KB with incremental sync + TOC sidecars."""

    def __init__(
        self,
        datalake_dir: str,
        embedding_client=None,
        embedding_deployment=None,
        llm_client=None,
        llm_deployment: str = None,
    ):
        self.datalake_dir = datalake_dir
        self.kb_dir = os.path.join(datalake_dir, KB_DIR_NAME)
        self.toc_dir = os.path.join(self.kb_dir, "toc")
        os.makedirs(self.toc_dir, exist_ok=True)
        self.chroma_dir = os.path.join(self.kb_dir, "chroma")
        os.makedirs(self.chroma_dir, exist_ok=True)
        self.manifest_path = os.path.join(self.kb_dir, MANIFEST_NAME)

        self.llm_client = llm_client
        self.llm_deployment = llm_deployment

        self._lock = threading.RLock()
        self._embed_fn = select_embedding_fn(embedding_client, embedding_deployment)
        self.client = chromadb.PersistentClient(path=self.chroma_dir)
        self.collection = self.client.get_or_create_collection(
            name="template_slides",
            metadata={"hnsw:space": "cosine"},
            embedding_function=self._embed_fn,
        )
        self.manifest = self._load_manifest()
        # TOC cache: filename -> parsed dict
        self._toc_cache: dict[str, dict] = {}

    # ── Manifest helpers ──

    def _load_manifest(self) -> dict:
        if os.path.isfile(self.manifest_path):
            try:
                with open(self.manifest_path, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception:
                pass
        return {}

    def _save_manifest(self):
        with open(self.manifest_path, "w", encoding="utf-8") as f:
            json.dump(self.manifest, f, indent=2)

    @staticmethod
    def _stat_signature(path: str) -> dict:
        st = os.stat(path)
        return {"mtime": st.st_mtime_ns, "size": st.st_size}

    @staticmethod
    def _template_key(filename: str) -> str:
        return hashlib.md5(filename.encode()).hexdigest()[:12]

    # ── Public sync ──

    def sync_index(self, log=print) -> dict:
        """Incremental scan: index new/changed templates, drop missing ones.

        Returns {added, updated, removed, skipped, total}.
        """
        with self._lock:
            on_disk = sorted(
                f for f in os.listdir(self.datalake_dir)
                if f.lower().endswith(".pptx") and not f.startswith("~$")
            )
            on_disk_set = set(on_disk)

            added = updated = skipped = removed = 0

            # Drop entries for templates that no longer exist
            for stale in list(self.manifest.keys()):
                if stale not in on_disk_set:
                    self._remove_template(stale)
                    removed += 1
                    log(f"  [KB] removed: {stale}")

            for fname in on_disk:
                path = os.path.join(self.datalake_dir, fname)
                sig = self._stat_signature(path)
                prev = self.manifest.get(fname)
                if prev and prev.get("mtime") == sig["mtime"] and prev.get("size") == sig["size"]:
                    skipped += 1
                    continue
                action = "added" if not prev else "updated"
                try:
                    self._index_template(path, fname, log=log)
                    self.manifest[fname] = sig
                    if action == "added":
                        added += 1
                    else:
                        updated += 1
                    log(f"  [KB] {action}: {fname}")
                except Exception as exc:
                    log(f"  [KB] FAILED {fname}: {exc}")

            self._save_manifest()
            total = self.collection.count()
            log(
                f"  [KB] sync done -- added={added} updated={updated} "
                f"removed={removed} skipped={skipped} total_slides_indexed={total}"
            )
            return {
                "added": added, "updated": updated, "removed": removed,
                "skipped": skipped, "total": total,
            }

    # ── Indexing internals ──

    def _remove_template(self, filename: str):
        try:
            self.collection.delete(where={"template_file": filename})
        except Exception:
            pass
        # Drop TOC sidecar
        slug = self._template_key(filename)
        toc_path = os.path.join(self.toc_dir, f"{slug}.json")
        if os.path.isfile(toc_path):
            try:
                os.remove(toc_path)
            except OSError:
                pass
        self.manifest.pop(filename, None)
        self._toc_cache.pop(filename, None)

    def _index_template(self, path: str, filename: str, log=print) -> int:
        # First, drop any stale entries from a prior version
        try:
            self.collection.delete(where={"template_file": filename})
        except Exception:
            pass

        prs = Presentation(path)
        template_name = os.path.splitext(filename)[0]

        all_text = []
        slides_meta = []  # raw extraction for sector / TOC

        for i, slide in enumerate(prs.slides):
            texts = []
            title = ""
            expanded = []
            for shape in slide.shapes:
                expanded.append(shape)
                if shape.shape_type == 6:  # group
                    try:
                        expanded.extend(shape.shapes)
                    except Exception:
                        pass
            for shape in expanded:
                if shape.has_text_frame:
                    t = shape.text_frame.text or ""
                    if t.strip():
                        texts.append(t.strip())
                    if "title" in (shape.name or "").lower() and t.strip() and not title:
                        title = t.strip()[:200]
                elif getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        for cell in row.cells:
                            ct = cell.text.strip()
                            if ct:
                                texts.append(ct)
            if not title and texts:
                title = texts[0][:200]
            content = "\n".join(texts)
            slides_meta.append({
                "index": i,
                "layout": slide.slide_layout.name if slide.slide_layout else "",
                "title": title,
                "content": content[:2000],
                "word_count": len(content.split()),
            })
            if content:
                all_text.append(content)

        sector = detect_sector("\n".join(all_text), filename)
        log(f"  [KB] {filename}: sector={sector}, slides={len(slides_meta)}")

        # Build TOC sidecar (LLM-assisted, cached on disk)
        toc = self._build_or_load_toc(filename, slides_meta, sector)
        purpose_by_index = {
            s["index"]: s.get("purpose", "other") for s in toc.get("slides", [])
        }

        # Chroma indexing — skip near-empty slides
        ids, documents, metadatas = [], [], []
        for sd in slides_meta:
            if sd["word_count"] < 5:
                continue
            section_type = purpose_by_index.get(sd["index"]) or classify_section(
                sd["title"], sd["content"]
            )
            # Normalize TOC's slide-purpose label back into our section_type vocab
            section_type = _normalize_section_type(section_type)
            doc_id = hashlib.md5(f"{template_name}_{sd['index']}".encode()).hexdigest()
            ids.append(doc_id)
            documents.append(sd["content"][:1500])
            metadatas.append({
                "template": template_name,
                "template_file": filename,
                "sector": sector,
                "section_type": section_type,
                "slide_index": sd["index"],
                "title": sd["title"][:200],
            })

        if ids:
            self.collection.add(ids=ids, documents=documents, metadatas=metadatas)
        return len(ids)

    def _build_or_load_toc(self, filename: str, slides_meta: list, sector: str) -> dict:
        slug = self._template_key(filename)
        toc_path = os.path.join(self.toc_dir, f"{slug}.json")
        if os.path.isfile(toc_path):
            try:
                with open(toc_path, "r", encoding="utf-8") as f:
                    cached = json.load(f)
                self._toc_cache[filename] = cached
                return cached
            except Exception:
                pass

        if self.llm_client and self.llm_deployment:
            toc = build_toc_with_llm(slides_meta, self.llm_client, self.llm_deployment)
        else:
            toc = {
                "sections": [{"name": "All slides", "purpose": "other",
                              "slide_range": [0, max(0, len(slides_meta) - 1)]}],
                "slides": [
                    {"index": s["index"], "title": s.get("title", ""),
                     "purpose": classify_section(s.get("title", ""), s.get("content", "")),
                     "summary": (s.get("content", "") or "")[:120]}
                    for s in slides_meta
                ],
            }

        toc["template"] = filename
        toc["sector"] = sector
        toc["slide_count"] = len(slides_meta)
        try:
            with open(toc_path, "w", encoding="utf-8") as f:
                json.dump(toc, f, indent=2)
        except Exception as exc:
            print(f"  [KB] could not persist TOC for {filename}: {exc}")

        self._toc_cache[filename] = toc
        return toc

    # ── Public lookups ──

    def get_toc(self, filename: str) -> Optional[dict]:
        """Return the per-template TOC dict, loading from disk if needed."""
        if filename in self._toc_cache:
            return self._toc_cache[filename]
        slug = self._template_key(filename)
        toc_path = os.path.join(self.toc_dir, f"{slug}.json")
        if os.path.isfile(toc_path):
            try:
                with open(toc_path, "r", encoding="utf-8") as f:
                    cached = json.load(f)
                self._toc_cache[filename] = cached
                return cached
            except Exception:
                return None
        return None

    def get_sector_tocs(self, sector: str, exclude_filename: str = None) -> list[dict]:
        """All TOCs for a given sector (useful for the structural-fill guardrail)."""
        out = []
        for fname in list(self.manifest.keys()):
            if exclude_filename and fname == exclude_filename:
                continue
            toc = self.get_toc(fname)
            if toc and (toc.get("sector") == sector or sector == "Any"):
                out.append(toc)
        return out

    def query(
        self, query_text: str, sector: str = None, section_type: str = None,
        n_results: int = 5,
    ) -> list[dict]:
        where_filters = {}
        if sector and section_type:
            where_filters = {"$and": [
                {"sector": {"$eq": sector}},
                {"section_type": {"$eq": section_type}},
            ]}
        elif sector:
            where_filters = {"sector": {"$eq": sector}}
        elif section_type:
            where_filters = {"section_type": {"$eq": section_type}}

        try:
            results = self.collection.query(
                query_texts=[query_text or " "],
                n_results=n_results,
                where=where_filters or None,
            )
        except Exception:
            results = self.collection.query(
                query_texts=[query_text or " "], n_results=n_results,
            )

        entries = []
        if results and results.get("documents") and results["documents"][0]:
            for doc, meta, dist in zip(
                results["documents"][0],
                results["metadatas"][0],
                results["distances"][0],
            ):
                entries.append({
                    "content": doc,
                    "template": meta.get("template", ""),
                    "template_file": meta.get("template_file", ""),
                    "sector": meta.get("sector", ""),
                    "section_type": meta.get("section_type", ""),
                    "title": meta.get("title", ""),
                    "relevance": round(1 - float(dist), 3),
                })
        return entries

    def get_sector_experience(self, sector: str, n_results: int = 8) -> list[dict]:
        return self.query(
            query_text=f"{sector} experience credentials case study project delivery",
            sector=sector, section_type="experience", n_results=n_results,
        )

    def get_sector_methodology(self, sector: str, n_results: int = 5) -> list[dict]:
        return self.query(
            query_text=f"{sector} methodology approach framework workstream",
            sector=sector, section_type="methodology", n_results=n_results,
        )

    def get_closest_sector_content(
        self, rfp_sector: str, section_type: str, query_text: str, n_results: int = 5,
    ) -> list[dict]:
        results = self.query(query_text, sector=rfp_sector, section_type=section_type, n_results=n_results)
        if not results:
            results = self.query(query_text, section_type=section_type, n_results=n_results)
        if not results:
            results = self.query(query_text, n_results=n_results)
        return results


def _normalize_section_type(t: str) -> str:
    """TOC LLM purposes -> our SECTION_TYPES keys."""
    if not t:
        return "other"
    t = t.lower()
    if t in SECTION_TYPES:
        return t
    aliases = {
        "title_slide": "cover", "section_header": "other",
        "diagram": "methodology", "table": "pricing",
        "cv": "team", "closing": "other", "divider": "other",
        "exec_summary": "understanding",
    }
    return aliases.get(t, "other")

import json
from pptx import Presentation


def extract_template_structure(pptx_path: str, client, deployment: str) -> list[dict]:
    prs = Presentation(pptx_path)
    raw_slides = _collect_slide_info(prs)
    classified = _classify_slides_with_llm(raw_slides, client, deployment)
    return classified


def _collect_slide_info(prs) -> list[dict]:
    slides_info = []
    for i, slide in enumerate(prs.slides):
        shapes_data = []
        for shape in slide.shapes:
            info = {
                "name": shape.name,
                "type": str(shape.shape_type),
                "has_text": shape.has_text_frame,
                "has_table": shape.has_table,
            }
            if shape.has_text_frame:
                text = shape.text_frame.text or ""
                info["text"] = text[:300]
                info["full_text"] = text
                info["para_count"] = len(shape.text_frame.paragraphs)
                info["word_count"] = len(text.split())
                # Detect bullet style
                has_bullets = False
                for p in shape.text_frame.paragraphs:
                    pPr = p._pPr
                    if pPr is not None:
                        buNone = pPr.find(
                            "{http://schemas.openxmlformats.org/drawingml/2006/main}buNone"
                        )
                        buChar = pPr.find(
                            "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar"
                        )
                        buAutoNum = pPr.find(
                            "{http://schemas.openxmlformats.org/drawingml/2006/main}buAutoNum"
                        )
                        if buChar is not None or buAutoNum is not None:
                            has_bullets = True
                            break
                info["has_bullets"] = has_bullets
            elif shape.has_table:
                tbl = shape.table
                info["table_size"] = f"{len(tbl.rows)}x{len(tbl.columns)}"
            shapes_data.append(info)

        slides_info.append(
            {
                "index": i,
                "layout": slide.slide_layout.name,
                "shapes": shapes_data,
            }
        )
    return slides_info


CLASSIFY_SYSTEM = """Analyze PowerPoint slides and classify each shape's role.

Return JSON:
{
  "slides": [
    {
      "index": 0,
      "section_heading": "Main topic/title of this slide (extracted from shapes)",
      "slide_purpose": "title_slide|section_header|content|agenda|appendix|closing|transition",
      "content_shapes": [
        {
          "name": "Shape Name",
          "role": "title|body|subtitle|decorative|label|footer|table_content",
          "replaceable": true
        }
      ]
    }
  ]
}

Classification rules:
- "title": The main heading shape (usually largest text, often named "Title *")
- "body": Main content area with substantial text (>10 words) meant to be replaced
- "subtitle": Secondary heading or subheading
- "decorative": Borders, logos, backgrounds, shapes with no/minimal text
- "label": Short labels (<5 words) that are part of visual design
- "footer": Page numbers, dates, confidentiality notices
- "table_content": Tables that could hold structured data
- Mark replaceable=true ONLY for title and body shapes with substantive text
- Shapes with empty text or <3 words are NOT replaceable
- Group shapes, OLE objects, connectors are NEVER replaceable
- Image shapes are NEVER replaceable"""


def _classify_slides_with_llm(
    slides_info: list[dict], client, deployment: str
) -> list[dict]:
    # Prepare a compact version for LLM (exclude full_text to save tokens)
    compact = []
    for s in slides_info:
        cs = {"index": s["index"], "layout": s["layout"], "shapes": []}
        for sh in s["shapes"]:
            csh = {"name": sh["name"], "type": sh["type"]}
            if sh.get("has_text"):
                csh["text_preview"] = sh.get("text", "")[:200]
                csh["word_count"] = sh.get("word_count", 0)
                csh["has_bullets"] = sh.get("has_bullets", False)
            if sh.get("has_table"):
                csh["table_size"] = sh.get("table_size")
            cs["shapes"].append(csh)
        compact.append(cs)

    batch_size = 8
    all_classified = []

    for start in range(0, len(compact), batch_size):
        batch = compact[start : start + batch_size]
        batch_json = json.dumps(batch, indent=2, default=str)
        print(f"    Classifying slides {start + 1}-{min(start + batch_size, len(compact))}...")

        resp = client.chat.completions.create(
            model=deployment,
            messages=[
                {"role": "system", "content": CLASSIFY_SYSTEM},
                {"role": "user", "content": f"Classify these slides:\n{batch_json}"},
            ],
            response_format={"type": "json_object"},
            temperature=0.05,
            max_tokens=8000,
        )
        result = json.loads(resp.choices[0].message.content)
        all_classified.extend(result.get("slides", []))

    # Enrich classified data with original full text and formatting info
    for cls_slide in all_classified:
        idx = cls_slide["index"]
        if idx < len(slides_info):
            orig = slides_info[idx]
            for cs in cls_slide.get("content_shapes", []):
                for orig_sh in orig["shapes"]:
                    if orig_sh["name"] == cs["name"]:
                        cs["_original_text"] = orig_sh.get("full_text", "")
                        cs["_word_count"] = orig_sh.get("word_count", 0)
                        cs["_para_count"] = orig_sh.get("para_count", 0)
                        cs["_has_bullets"] = orig_sh.get("has_bullets", False)
                        break

    return all_classified

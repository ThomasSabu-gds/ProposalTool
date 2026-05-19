"""
Verifier Agent: post-generation quality pass.

1. Content Verification  — LLM checks each slide for hallucination against KB/RFP
2. Layout Enforcement    — Python checks text overflow, title length, shape bounds
3. Slide Curation        — LLM reviews full deck, deletes redundant slides, keeps 1 CV
"""

import json
import re
import copy
from pptx import Presentation
from pptx.util import Emu, Pt
from lxml import etree

from knowledge_base import classify_section

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

_CTRL_CHAR_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


# ═══════════════════════════════════════════════════════════════════════════
# MAIN ENTRY POINT
# ═══════════════════════════════════════════════════════════════════════════

def verify_and_fix(
    pptx_path: str,
    client, deployment: str,
    rfp_summary_text: str,
    kb_entries: list,
    client_name: str,
    rfp_sector: str,
) -> dict:
    """Run all verification passes on the generated PPTX. Modifies in place."""
    print("  [Verify] Starting verification pass...")
    prs = Presentation(pptx_path)
    report = {"layout_fixes": 0, "content_fixes": 0, "slides_deleted": 0, "slides_total": len(prs.slides)}

    # ── Pass 1: Layout enforcement (fast, no LLM) ──
    print("  [Verify] Pass 1: Layout enforcement...")
    layout_fixes = _enforce_layout(prs)
    report["layout_fixes"] = layout_fixes
    print(f"  [Verify]   Fixed {layout_fixes} overflow/layout issues")

    # ── Pass 2: Content verification (LLM per-slide batch) ──
    print("  [Verify] Pass 2: Content verification...")
    content_fixes = _verify_content(prs, client, deployment, rfp_summary_text, kb_entries, client_name, rfp_sector)
    report["content_fixes"] = content_fixes
    print(f"  [Verify]   Fixed {content_fixes} content issues")

    # ── Pass 3: Slide curation (LLM reviews full deck) ──
    print("  [Verify] Pass 3: Slide curation...")
    deleted = _curate_slides(prs, client, deployment, rfp_summary_text, rfp_sector)
    report["slides_deleted"] = deleted
    report["slides_final"] = len(prs.slides)
    print(f"  [Verify]   Deleted {deleted} slides, {len(prs.slides)} remaining")

    prs.save(pptx_path)
    print(f"  [Verify] Verification complete: {report}")
    return report


# ═══════════════════════════════════════════════════════════════════════════
# PASS 1: LAYOUT ENFORCEMENT
# ═══════════════════════════════════════════════════════════════════════════

# Approximate chars per inch at common font sizes (conservative estimates)
_CHARS_PER_INCH = {
    8: 18, 9: 16, 10: 14, 11: 13, 12: 12, 14: 10,
    16: 9, 18: 8, 20: 7, 22: 6, 24: 5.5, 28: 4.5,
    32: 4, 36: 3.5, 40: 3, 44: 2.5, 48: 2.2,
}

# Approximate line height in inches at common font sizes
_LINE_HEIGHT_INCHES = {
    8: 0.15, 9: 0.17, 10: 0.19, 11: 0.21, 12: 0.23, 14: 0.27,
    16: 0.31, 18: 0.35, 20: 0.39, 22: 0.43, 24: 0.47, 28: 0.55,
    32: 0.63, 36: 0.71, 40: 0.78, 44: 0.86, 48: 0.94,
}


def _is_section_divider_slide(slide) -> bool:
    """Recognize a chapter-marker slide: ONE 1-2 digit number plus ONE short
    section name (1-7 words). Multiple numbers / multiple labels indicate a
    TOC or a step diagram, not a section divider — those must be left alone.
    We tolerate AT MOST ONE long paragraph (the leak we are trying to undo)."""
    number_count = 0
    title_count = 0
    long_text_shapes = 0
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = (shape.text_frame.text or "").strip()
        if not text:
            continue
        wc = len(text.split())
        cleaned = text.replace(".", "").strip()
        if 1 <= len(text) <= 4 and cleaned.isdigit():
            number_count += 1
            continue
        if 2 <= wc <= 7 and not text.endswith("."):
            title_count += 1
        elif wc == 1 and len(text) >= 6 and not text.endswith("."):
            title_count += 1
        elif wc >= 15:
            long_text_shapes += 1
    return (
        number_count == 1
        and title_count == 1
        and long_text_shapes <= 1
    )


def _enforce_layout(prs) -> int:
    fixes = 0
    for slide in prs.slides:
        # Safety net: if this slide is a section divider (e.g. "01 / Why Us"),
        # blank out any body shape that holds a long paragraph. The slide-agent
        # is told to skip these but a defense-in-depth wipe protects the design
        # when older runs or upstream changes slip a paragraph through.
        if _is_section_divider_slide(slide):
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                text = shape.text_frame.text or ""
                wc = len(text.split())
                if wc >= 12:
                    _replace_text_in_frame(shape.text_frame, "")
                    fixes += 1
            continue

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            if not shape.text_frame.text.strip():
                continue

            w_inches = shape.width / 914400 if shape.width else 0
            h_inches = shape.height / 914400 if shape.height else 0
            if w_inches <= 0 or h_inches <= 0:
                continue

            try:
                shape.text_frame.word_wrap = True
            except Exception:
                pass

            font_size = _get_dominant_font_size(shape.text_frame)
            if not font_size:
                continue

            name_l = shape.name.lower()
            is_title = "title" in name_l
            is_heading = is_title or any(t in name_l for t in ("heading", "subtitle", "header"))

            # First pass for every shape: try to shrink the font so the existing
            # text fits inside the physical box. This is the gentlest fix and
            # preserves meaning.
            shrunk_to = _shrink_font_to_fit(
                shape.text_frame, w_inches, h_inches, font_size,
                min_pt=14 if is_title else 8,
            )
            effective_font = shrunk_to or font_size
            if shrunk_to and shrunk_to != font_size:
                fixes += 1

            if is_heading:
                # Even after shrinking a title may still spill onto > 2 lines
                # if the LLM wrote a long phrase. Trim conservatively.
                if _fix_title_overflow(shape.text_frame, w_inches, h_inches, effective_font):
                    fixes += 1
                continue

            # Body text: trim if it still overflows after shrinking
            if _fix_body_overflow(shape.text_frame, w_inches, h_inches, effective_font):
                fixes += 1

    return fixes


def _shrink_font_to_fit(text_frame, w_inches, h_inches, current_font, min_pt=8):
    """Iteratively reduce the font size on every run until the text fits the
    shape's box (within ~5% slack). Returns the new size if any change was
    made, else None. Never goes below min_pt."""
    text = text_frame.text or ""
    if not text.strip():
        return None

    new_size = current_font
    while new_size > min_pt:
        max_cpl = _estimate_max_chars_per_line(w_inches, new_size)
        max_ln = _estimate_max_lines(h_inches, new_size)
        capacity = max_cpl * max_ln
        # Estimate wrapped line count
        lines = text.split("\n")
        est_lines = 0
        for line in lines:
            if not line.strip():
                est_lines += 1
            else:
                est_lines += max(1, (len(line) + max_cpl - 1) // max_cpl)
        if est_lines <= max_ln and len(text) <= capacity * 1.05:
            break
        new_size -= 1

    if new_size == current_font:
        return None

    for para in text_frame.paragraphs:
        for run in para.runs:
            try:
                run.font.size = Pt(new_size)
            except Exception:
                pass
    return new_size


def _get_dominant_font_size(text_frame) -> float:
    """Get the most common font size in a text frame (in points)."""
    sizes = []
    for para in text_frame.paragraphs:
        for run in para.runs:
            if run.font.size:
                sizes.append(round(run.font.size / 12700))
    if not sizes:
        return 12  # default
    # Return most common
    from collections import Counter
    return Counter(sizes).most_common(1)[0][0]


def _estimate_max_chars_per_line(w_inches: float, font_size: float) -> int:
    """Estimate max characters per line given shape width and font size."""
    fs_key = min(_CHARS_PER_INCH.keys(), key=lambda k: abs(k - font_size))
    cpi = _CHARS_PER_INCH[fs_key]
    return max(10, int(w_inches * cpi * 0.85))  # 85% to account for margins


def _estimate_max_lines(h_inches: float, font_size: float) -> int:
    """Estimate max lines that fit in a shape given its height and font size."""
    fs_key = min(_LINE_HEIGHT_INCHES.keys(), key=lambda k: abs(k - font_size))
    lh = _LINE_HEIGHT_INCHES[fs_key]
    return max(1, int(h_inches / lh * 0.85))  # 85% for padding


def _fix_title_overflow(text_frame, w_inches, h_inches, font_size) -> bool:
    """Ensure title fits in 1-2 lines. Truncate if needed."""
    max_chars = _estimate_max_chars_per_line(w_inches, font_size)
    max_lines = min(2, _estimate_max_lines(h_inches, font_size))
    max_total_chars = max_chars * max_lines

    full_text = text_frame.text.replace("\n", " ").strip()
    if len(full_text) <= max_total_chars:
        return False

    # Truncate intelligently - cut at word boundary
    truncated = full_text[:max_total_chars].rsplit(" ", 1)[0]
    if len(truncated) < len(full_text) * 0.5:
        truncated = full_text[:max_total_chars]

    _replace_text_in_frame(text_frame, truncated)
    return True


def _fix_body_overflow(text_frame, w_inches, h_inches, font_size) -> bool:
    """Trim body text when it still overflows after font shrink (>=1.2x capacity)."""
    max_chars_per_line = _estimate_max_chars_per_line(w_inches, font_size)
    max_lines = _estimate_max_lines(h_inches, font_size)

    lines = text_frame.text.split("\n")
    estimated_lines = 0
    for line in lines:
        if not line.strip():
            estimated_lines += 1
        else:
            wrap_lines = max(1, (len(line) + max_chars_per_line - 1) // max_chars_per_line)
            estimated_lines += wrap_lines

    overflow_threshold = max(max_lines + 1, int(max_lines * 1.2))
    if estimated_lines <= overflow_threshold:
        return False

    # Trim to fit cleanly (1.0x max_lines). Shrink-to-fit already happened.
    trim_budget = max(max_lines, int(max_lines * 1.0))
    kept_lines = []
    line_count = 0
    for line in lines:
        if not line.strip():
            add = 1
        else:
            add = max(1, (len(line) + max_chars_per_line - 1) // max_chars_per_line)
        if line_count + add > trim_budget:
            if line_count < trim_budget:
                remaining_chars = (trim_budget - line_count) * max_chars_per_line
                kept_lines.append(line[:remaining_chars].rsplit(" ", 1)[0])
            break
        kept_lines.append(line)
        line_count += add

    new_text = "\n".join(kept_lines)
    if new_text != text_frame.text:
        _replace_text_in_frame(text_frame, new_text)
        return True
    return False


# ═══════════════════════════════════════════════════════════════════════════
# PASS 2: CONTENT VERIFICATION (anti-hallucination)
# ═══════════════════════════════════════════════════════════════════════════

VERIFY_CONTENT_SYSTEM = """You are a strict proposal content reviewer. You check for HALLUCINATION and INACCURACY.

You receive a slide's content, the RFP summary, and the FULL TEXT of every knowledge base entry that backs this proposal.
Treat the RFP and the KB entries together as the ONLY allowed source of named facts.

YOU MUST FLAG (replace with a generic, source-backed version):
1. HALLUCINATED CLIENT ENGAGEMENT: any sentence that claims or implies the firm has worked with the RFP client. The RFP client is PROSPECTIVE — we have NOT worked with them.
2. FABRICATED NAMED ENTITY: any specific organization name, brand, project name, acronym (e.g. "LEHS", "ACME Health"), program name, or proper-noun client reference that does NOT appear verbatim in the KB entries or RFP summary. If you cannot find that exact name in the provided sources, it is fabricated.
3. FABRICATED FACTS: numbers, percentages, dates, durations, dollar amounts, team sizes, KPIs, or statistics that do not appear in the RFP or KB sources.
4. WRONG SECTOR: case studies framed in an industry that contradicts the RFP sector when the KB does not support that framing.
5. MADE-UP PROJECT TITLES: invented case study titles or project headers not present in the KB.

FIX RULE:
- When you flag a hallucinated named entity or fabricated fact, rewrite the affected sentence by replacing the bad name/number with a sourced generalization
  (e.g. "a leading healthcare provider", "a national health authority", "multiple multi-year engagements") and keep the rest of the surrounding text intact.
- Preserve headings, section titles, and the overall meaning. Do NOT delete entire bullets or paragraphs just to remove a name — rewrite minimally.
- Never invent a replacement entity. Only use phrasing supported by KB/RFP.

Return JSON:
{
  "issues": [
    {"shape": "Shape Name", "problem": "brief description", "fix": "corrected text for the WHOLE shape (preserve unaffected lines verbatim)"}
  ]
}

If no issues found, return: {"issues": []}

DO NOT flag:
- Generic advisory language ("a leading entity", "a major authority", "global firm", "Big Four firm")
- Sector-adapted paraphrasing of KB content that keeps real facts
- The firm's own name ("EY", "Ernst & Young") used as the proposing firm
- Generic methodology, framework, or workstream names that are clearly descriptive rather than proper nouns"""


def _verify_content(prs, client, deployment, rfp_summary_text, kb_entries, client_name, rfp_sector) -> int:
    fixes = 0

    enriched_slides = []
    for i, slide in enumerate(prs.slides):
        text_shapes = []
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text.strip() and len(sh.text) > 20:
                text_shapes.append({"name": sh.name, "text": sh.text[:600]})
        if text_shapes:
            section_type = classify_section(
                text_shapes[0]["text"][:50] if text_shapes else "",
                " ".join(s["text"] for s in text_shapes),
            )
            enriched_slides.append({
                "slide_index": i,
                "section_type": section_type,
                "shapes": text_shapes,
            })

    kb_text = ""
    if kb_entries:
        for j, entry in enumerate(kb_entries[:8]):
            kb_text += (
                f"\nKB{j+1} [template={entry.get('template','')}|sector={entry.get('sector','')}|"
                f"type={entry.get('section_type','')}]:\n{entry.get('content','')[:900]}\n"
            )
    kb_payload = kb_text[:7000] if kb_text else "(no KB entries available for this sector)"

    # Verify slide-by-slide so the model has full visibility per slide
    for s_data in enriched_slides:
        slide_payload = json.dumps({
            "slide": s_data["slide_index"] + 1,
            "type": s_data["section_type"],
            "shapes": s_data["shapes"],
        }, indent=1)[:5000]

        try:
            resp = client.chat.completions.create(
                model=deployment,
                messages=[
                    {"role": "system", "content": VERIFY_CONTENT_SYSTEM},
                    {"role": "user", "content": (
                        f"RFP CLIENT NAME (PROSPECTIVE — we have NOT worked with them): {client_name}\n"
                        f"RFP SECTOR: {rfp_sector}\n\n"
                        f"RFP SUMMARY (source of allowed RFP facts):\n{rfp_summary_text[:3000]}\n\n"
                        f"KNOWLEDGE BASE (source of allowed firm facts and named clients/projects):\n{kb_payload}\n\n"
                        f"SLIDE TO VERIFY:\n{slide_payload}\n\n"
                        "Cross-check every proper noun, organization name, acronym, statistic, and dollar/duration/team-size figure in the slide against the RFP summary and KB entries above. "
                        "If you cannot find a name or number in those sources, treat it as fabricated and rewrite that sentence using a sourced generalization. "
                        "Preserve all unaffected text verbatim in the 'fix' field."
                    )},
                ],
                response_format={"type": "json_object"},
                temperature=0.05,
                max_tokens=4096,
            )
            result = json.loads(resp.choices[0].message.content)
            issues = result.get("issues", [])

            for issue in issues:
                shape_name = issue.get("shape", "")
                fix_text = issue.get("fix", "")
                if not shape_name or not fix_text:
                    continue
                slide = prs.slides[s_data["slide_index"]]
                for sh in slide.shapes:
                    if sh.has_text_frame and sh.name == shape_name:
                        _replace_text_in_frame(sh.text_frame, fix_text)
                        fixes += 1
                        print(f"    [Verify] Fixed: slide {s_data['slide_index']+1} '{shape_name}': {issue.get('problem','')}")
                        break

        except Exception as e:
            print(f"    [Verify] Slide {s_data['slide_index']+1} verify error: {e}")

    return fixes


# ═══════════════════════════════════════════════════════════════════════════
# PASS 3: SLIDE CURATION (delete redundant, keep 1 CV)
# ═══════════════════════════════════════════════════════════════════════════

CURATE_SYSTEM = """You are reviewing a proposal deck for a client presentation. Your job is to decide which slides to DELETE.

Rules:
1. DELETE slides that are completely redundant (nearly identical content to another slide)
2. DELETE slides that are irrelevant to the RFP's sector and project
3. DELETE slides that are mostly empty or have ONLY generic filler text and no useful structure
4. KEEP exactly 1 populated team/CV slide; if there are MANY blank/template CV pages, KEEP 1-2 of them (so the user has a CV template they can fill in later) and DELETE the rest. Never delete ALL blank CV slides.
5. KEEP all slides with substantive RFP-specific content
6. KEEP cover, cover letter, TOC, and closing slides
7. KEEP methodology and approach slides
8. KEEP experience slides (but remove obvious duplicates)
9. KEEP section-divider / chapter-header slides (slides whose only content is a section number like "01" / "02" plus a short section name like "Why Us" / "Methodology"). These are visual chapter markers and must be preserved as-is.

Return JSON:
{
  "delete": [slide_number_1, slide_number_2],
  "reason": {"slide_number": "reason for deletion"}
}

If no slides should be deleted, return: {"delete": [], "reason": {}}
Be CONSERVATIVE — only delete slides you are SURE are redundant or irrelevant."""


def _curate_slides(prs, client, deployment, rfp_summary_text, rfp_sector) -> int:
    # Build slide outline
    outline = []
    blank_cv_idxs = []  # zero-based indices of empty CV-style template pages
    for i, slide in enumerate(prs.slides):
        texts = []
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text.strip():
                texts.append(sh.text[:100].replace("\n", " "))
        section_type = classify_section(
            texts[0][:50] if texts else "",
            " ".join(texts),
        )
        wc = sum(len(t.split()) for t in texts)
        layout_name = (slide.slide_layout.name or "").lower()
        # Empty CV template page: blank/very sparse + uses a CV-style layout.
        if wc <= 6 and (
            "cv" in layout_name
            or "1/3 column" in layout_name
            or "1_1/3" in layout_name
        ):
            blank_cv_idxs.append(i)
        outline.append({
            "slide": i + 1,
            "type": section_type,
            "content_preview": " | ".join(texts[:3])[:200] if texts else "(empty)",
            "word_count": wc,
        })

    try:
        resp = client.chat.completions.create(
            model=deployment,
            messages=[
                {"role": "system", "content": CURATE_SYSTEM},
                {"role": "user", "content": (
                    f"RFP SECTOR: {rfp_sector}\n\n"
                    f"RFP SUMMARY:\n{rfp_summary_text[:1500]}\n\n"
                    f"SLIDE OUTLINE ({len(outline)} slides):\n"
                    f"{json.dumps(outline, indent=1)}"
                )},
            ],
            response_format={"type": "json_object"},
            temperature=0.05,
            max_tokens=2048,
        )
        result = json.loads(resp.choices[0].message.content)
        to_delete = result.get("delete", [])
        reasons = result.get("reason", {})

        if not to_delete:
            return 0

        # Convert to 0-based indices and sort descending (delete from end first)
        indices = sorted([int(n) - 1 for n in to_delete if 0 < int(n) <= len(prs.slides)], reverse=True)

        # Protect 1-2 blank CV/team template pages so the user always has a
        # CV slide they can fill in later. We keep the FIRST two such pages
        # (if present) and remove them from the deletion list.
        protected_blanks = set(blank_cv_idxs[:2])
        if protected_blanks:
            removed = [i for i in indices if i in protected_blanks]
            if removed:
                indices = [i for i in indices if i not in protected_blanks]
                for i in removed:
                    print(f"    [Verify] Preserving slide {i+1}: blank CV template kept for user")

        deleted = 0
        for idx in indices:
            reason = reasons.get(str(idx + 1), "redundant")
            print(f"    [Verify] Deleting slide {idx+1}: {reason}")
            _delete_slide(prs, idx)
            deleted += 1

        return deleted

    except Exception as e:
        print(f"    [Verify] Curation error: {e}")
        return 0


def _delete_slide(prs, slide_index: int):
    """Delete a slide by index from the presentation."""
    if slide_index >= len(prs.slides):
        return
    rId = prs.slides._sldIdLst[slide_index].get(f"{{{NS_R}}}id")
    if rId is None:
        # Try the 'id' attribute approach
        sldId = prs.slides._sldIdLst[slide_index]
        # Remove from sldIdLst
        prs.slides._sldIdLst.remove(sldId)
        return

    sldId = prs.slides._sldIdLst[slide_index]
    prs.slides._sldIdLst.remove(sldId)
    # Remove the relationship
    try:
        prs.part.drop_rel(rId)
    except Exception:
        pass


# ═══════════════════════════════════════════════════════════════════════════
# TEXT REPLACEMENT HELPERS
# ═══════════════════════════════════════════════════════════════════════════

def _replace_text_in_frame(text_frame, new_text: str):
    """Replace all text in a text frame preserving formatting of first paragraph."""
    txBody = text_frame._txBody
    a_p = f"{{{NS_A}}}p"
    a_r = f"{{{NS_A}}}r"
    a_t = f"{{{NS_A}}}t"
    a_rPr = f"{{{NS_A}}}rPr"
    a_pPr = f"{{{NS_A}}}pPr"
    a_endParaRPr = f"{{{NS_A}}}endParaRPr"

    existing_paras = [child for child in txBody if child.tag == a_p]
    if not existing_paras:
        return

    # Get formatting from first paragraph with content
    fmt_pPr = None
    fmt_rPr = None
    fmt_endParaRPr = None
    for p in existing_paras:
        pPr = p.find(a_pPr)
        runs = p.findall(a_r)
        if runs:
            rPr = runs[0].find(a_rPr)
            if rPr is not None:
                fmt_rPr = copy.deepcopy(rPr)
            if pPr is not None:
                fmt_pPr = copy.deepcopy(pPr)
            endPr = p.find(a_endParaRPr)
            if endPr is not None:
                fmt_endParaRPr = copy.deepcopy(endPr)
            break

    # Remove old paragraphs
    for p in existing_paras:
        txBody.remove(p)

    # Build new paragraphs
    new_lines = new_text.split("\n") if new_text else [""]
    for line in new_lines:
        p_elem = etree.SubElement(txBody, a_p)
        if fmt_pPr is not None:
            p_elem.insert(0, copy.deepcopy(fmt_pPr))
        r_elem = etree.SubElement(p_elem, a_r)
        if fmt_rPr is not None:
            r_elem.insert(0, copy.deepcopy(fmt_rPr))
        t_elem = etree.SubElement(r_elem, a_t)
        t_elem.text = _sanitize(line)
        t_elem.set(XML_SPACE, "preserve")
        if fmt_endParaRPr is not None:
            p_elem.append(copy.deepcopy(fmt_endParaRPr))


def _sanitize(text: str) -> str:
    return _CTRL_CHAR_RE.sub("", text) if text else ""

"""Dump per-slide text from the produced PPTX so we can see what landed."""
import os, sys
from pptx import Presentation

path = sys.argv[1] if len(sys.argv) > 1 else None
if not path:
    runs = r"C:\Users\MP946KB\WORKDIR\Proposal-assist\ProposalAssistant\runs"
    cands = sorted(
        (os.path.join(runs, f) for f in os.listdir(runs)
         if f.endswith(".pptx") and not f.startswith("~$")),
        key=os.path.getmtime, reverse=True,
    )
    path = cands[0]

print(f"FILE: {path}\n")
prs = Presentation(path)
for i, slide in enumerate(prs.slides):
    print(f"=== Slide {i+1} (layout={slide.slide_layout.name}) ===")
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text.strip()
        if not txt:
            continue
        first = txt.replace("\n", " | ")[:200]
        print(f"  [{sh.name}] {first}")
    print()
